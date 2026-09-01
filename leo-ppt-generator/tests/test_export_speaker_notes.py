#!/usr/bin/env python3
"""export_speaker_notes.py 单测：PPTX notes 导出 / 母版 speaker_script 导出 /
缺备注诚实列出 / 确定性 / 互斥与缺参 exit 2。"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "export_speaker_notes.py"

MASTER = """# 逐页母版 v1 — 测试 deck

## P1 封面
- speaker_script：各位好，今天讲测试方案。（停顿）
- engineering：封面版式。

## P2 方法
- speaker_script：方法一句话：先 A 后 B。

## P3 收束
- engineering：无口播稿的收束页。
"""


def _run(args, expect_ok=True):
    result = subprocess.run(
        [sys.executable, str(SCRIPT)] + [str(a) for a in args],
        capture_output=True, text=True,
    )
    if expect_ok and result.returncode != 0:
        raise AssertionError(f"expected exit 0, got {result.returncode}: {result.stderr}")
    return result


def _fixture_pptx(path: Path):
    presentation = Presentation()
    slide1 = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide1.notes_slide.notes_text_frame.text = "第一页口播：开场。"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide2.notes_slide.notes_text_frame.text = "第二页口播：收束。"
    presentation.save(str(path))


class PptxExportTest(unittest.TestCase):
    def test_exports_notes_in_order_with_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            pptx = Path(tmp) / "deck.pptx"
            _fixture_pptx(pptx)
            out = _run(["--pptx", pptx]).stdout
        self.assertIn("第 1 页", out)
        self.assertIn("第一页口播：开场。", out)
        self.assertIn("第二页口播：收束。", out)
        self.assertIn("共 2 页；有备注 2 页，缺备注 0 页", out)

    def test_missing_notes_listed_honestly(self):
        with tempfile.TemporaryDirectory() as tmp:
            pptx = Path(tmp) / "deck.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])  # no notes
            presentation.save(str(pptx))
            out = _run(["--pptx", pptx]).stdout
        self.assertIn("（无备注）", out)
        self.assertIn("缺备注 1 页", out)


class MasterExportTest(unittest.TestCase):
    def test_extracts_speaker_script_per_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "deck-master-v1.md"
            master.write_text(MASTER, encoding="utf-8")
            out = _run(["--master", master]).stdout
        self.assertIn("第 1 页：封面", out)
        self.assertIn("今天讲测试方案", out)
        self.assertIn("第 3 页：收束", out)
        self.assertIn("（无备注）", out)
        self.assertIn("共 3 页；有备注 2 页，缺备注 1 页", out)

    def test_master_without_page_headings_exits_2(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "bad.md"
            master.write_text("没有分页的母版\n", encoding="utf-8")
            result = _run(["--master", master], expect_ok=False)
        self.assertEqual(result.returncode, 2)


class ContractTest(unittest.TestCase):
    def test_no_source_exits_2(self):
        result = subprocess.run([sys.executable, str(SCRIPT)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)

    def test_both_sources_exit_2(self):
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "--pptx", "a.pptx", "--master", "b.md"],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 2)

    def test_deterministic_output_and_out_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "m.md"
            master.write_text(MASTER, encoding="utf-8")
            out1 = _run(["--master", master]).stdout
            out2 = _run(["--master", master]).stdout
            self.assertEqual(out1, out2)
            target = Path(tmp) / "script.md"
            _run(["--master", master, "--out", target])
            self.assertEqual(target.read_text(encoding="utf-8"), out1)


class ProseCheckTest(unittest.TestCase):
    """--prose-check：R-16 讲稿纪律 WARN 汇总（stderr、不阻断导出、不改退出码）。"""

    def test_prose_check_prints_warn_summary_without_blocking(self):
        master = """# 逐页母版 v1 — prose-check 测试

## P1 方法
- speaker_script：接下来我将为大家介绍方法。这一页的讲稿句子故意写得很长很长并且明显超过四十个字的上限以便触发长句警告的测试用例现在应该足够长了。

## P2 收束
- speaker_script：先说结论。
"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck-master-v1.md"
            path.write_text(master, encoding="utf-8")
            result = _run(["--master", path, "--prose-check"])
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("PROSE-WARN", result.stderr)
        self.assertIn("第 1 页", result.stderr)
        self.assertIn("讲稿套话「接下来我将」", result.stderr)
        # export itself still happens on stdout
        self.assertIn("# 讲稿（来源：deck master", result.stdout)
        self.assertIn("共 2 页；有备注 2 页，缺备注 0 页", result.stdout)

    def test_prose_check_clean_script_is_silent(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "deck-master-v1.md"
            master.write_text(MASTER, encoding="utf-8")
            result = _run(["--master", master, "--prose-check"])
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("PROSE-WARN", result.stderr)
        self.assertNotIn("PROSE-CHECK", result.stderr)

    def test_prose_check_with_out_file_keeps_export_intact(self):
        master = """## P1 方法
- speaker_script：综上所述，先看结论。
"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck-master-v1.md"
            path.write_text(master, encoding="utf-8")
            target = Path(tmp) / "script.md"
            result = _run(["--master", path, "--prose-check", "--out", target])
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("PROSE-WARN", result.stderr)
            written = target.read_text(encoding="utf-8")
        self.assertIn("# 讲稿（来源：deck master", written)
        self.assertNotIn("PROSE-WARN", written)


if __name__ == "__main__":
    unittest.main()
