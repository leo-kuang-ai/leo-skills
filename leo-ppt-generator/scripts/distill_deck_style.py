#!/usr/bin/env python3
"""distill_deck_style.py — 标杆 deck 蒸馏档案生成器（R-54，python-pptx）。

入口语义（与 extract_pptx_theme.py 同管道）：**仅接受用户已确认可信的
PPTX**——Office Trust Gate（Gate 0）语义由调用方负责，本脚本不做信任判定
（见 references/input-routing.md、references/deck-distillation.md）。

从可信标杆 deck 确定性提取五层蒸馏档案（借鉴 blogger-distiller 五层结构，
出处粒度适配为页码）：

1. 论证模式观察 —— 开篇/收束角色、图表承载论证占比、标题断言形态占比；
2. 页面节奏 —— 每页版式角色序列（确定性规则推断）；
3. 版式偏好 —— 各角色页的文本框数量/图表频率分布；
4. 标题风格 —— 标题样本、长度分布、句末标点形态；
5. 样本局限声明 —— 覆盖率如实声明（页数、标题覆盖率、未采样维度）。

**防编造硬规则**（与 blogger-distiller「所有公式必须有原始笔记作为来源」
同构）：每条观察必须携带页码出处（``(p2, p5)``）；``pages`` 为空的观察在
渲染层被无条件丢弃——无出处的概括不允许出现在档案中，防止把通用版式
包装成「该 deck 独有」。

观察验证状态（E6-4「≥2 样本验证」的确定性对应）：支撑页 ≥2 → 「多页
复现」；=1 → 「单页观察，反演时归『需确认』组」。

确定性：无时间戳、无随机；直方图按 (-count, key) 排序，输出可复现。

版式角色规则（简单确定性映射，非风格判断）：

    charts>=1 → chart；tables>=1 → table；pictures>=3 → image-grid；
    pictures>=1 → text-image；texts>=4 → dense-text；texts>=1 →
    sparse-text；否则 blank。

字号分档阈值（观察分档，非设计规范——投屏字号下限另见
references/visual-qa.md 第四节判据表）：≥28pt=h1；18–28=h2；14–18=body；
<14=caption；未显式设置=inherit。

色板来源限于显式 RGB：文本 run 颜色 + 实心填充前景色；主题 scheme 色引用
单独计数（不展开为 HEX，防止把主题槽位臆测成具体色值）。

档案落盘约定（脚本不强制建目录）：``${LEO_PPT_HOME}/decks-styles/<名>.md``，
与 brands/profiles 同构的可点名资产（风格优先序：点名 > 参考图 > 推荐，
见 references/deck-distillation.md）。

用法::

    python3 scripts/distill_deck_style.py <可信.pptx> [--json|-] [--out 档案.md]

退出码：0 成功；2 文件/用法错误（缺参、文件不存在、非 PPTX、0 页空 deck）。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import MSO_COLOR_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

# Observation bucket thresholds; deliberately independent from the delivery
# spec in visual-qa.md §4 (those are authoring floors, these are read buckets).
FONT_TIER_BOUNDS = ((28.0, "h1"), (18.0, "h2"), (14.0, "body"))
FONT_TIER_FALLBACK = "caption"
FONT_TIER_UNSPECIFIED = "inherit"
TIER_ORDER = ("h1", "h2", "body", "caption", FONT_TIER_UNSPECIFIED)

PALETTE_TOP_N = 8

EXIT_USAGE = 2

# Support labels for the >=2-sample verification convention (E6-4).
LABEL_MULTI_PAGE = "多页复现"
LABEL_SINGLE_PAGE = "单页观察，反演时归「需确认」组"


def _font_tier(pt: float | None) -> str:
    if pt is None:
        return FONT_TIER_UNSPECIFIED
    for bound, tier in FONT_TIER_BOUNDS:
        if pt >= bound:
            return tier
    return FONT_TIER_FALLBACK


def _page_role(charts: int, tables: int, pictures: int, texts: int) -> str:
    # Deterministic layout-role mapping; ordering is the contract under test.
    if charts >= 1:
        return "chart"
    if tables >= 1:
        return "table"
    if pictures >= 3:
        return "image-grid"
    if pictures >= 1:
        return "text-image"
    if texts >= 4:
        return "dense-text"
    if texts >= 1:
        return "sparse-text"
    return "blank"


def _bump(counter: dict, key: str, page: int) -> None:
    entry = counter.setdefault(key, {"count": 0, "pages": []})
    entry["count"] += 1
    if page not in entry["pages"]:
        entry["pages"].append(page)


def _collect_shape(shape, page: int, stats: dict) -> None:
    # Text runs: font size buckets + explicit RGB colors + scheme refs.
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                size = run.font.size
                pt = size.pt if size is not None else None
                _bump(stats["font_tiers"], _font_tier(pt), page)
                if pt is not None:
                    _bump(stats["font_pts"], _repr_pt(pt), page)
                color = run.font.color
                try:
                    if color.type == MSO_COLOR_TYPE.RGB:
                        _bump(stats["palette"], str(color.rgb), page)
                    elif color.type == MSO_COLOR_TYPE.SCHEME:
                        _bump(stats["scheme_refs"], str(color.theme_color), page)
                except (TypeError, ValueError):
                    pass  # inherited/unresolvable color: not explicit evidence
    # Solid fills: explicit RGB only. GraphicFrame (chart/table) has no fill.
    fill = getattr(shape, "fill", None)
    if fill is None:
        return
    try:
        if fill.type == MSO_FILL_TYPE.SOLID:
            fore = fill.fore_color
            if fore.type == MSO_COLOR_TYPE.RGB:
                _bump(stats["palette"], str(fore.rgb), page)
    except (TypeError, ValueError, RuntimeError, AttributeError):
        pass


def _repr_pt(pt: float) -> str:
    return f"{pt:g}pt"


def extract(pptx_path: Path) -> dict:
    """Read deck into a deterministic, JSON-serialisable evidence dict."""
    prs = Presentation(str(pptx_path))
    slides: list[dict] = []
    stats = {
        "font_tiers": {},
        "font_pts": {},
        "palette": {},
        "scheme_refs": {},
    }
    for index, slide in enumerate(prs.slides, start=1):
        charts = tables = pictures = texts = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts += 1
            _collect_shape(shape, index, stats)
        title_shape = slide.shapes.title
        title = None
        if title_shape is not None and title_shape.text_frame.text.strip():
            title = title_shape.text_frame.text.strip()
        slides.append({
            "page": index,
            "role": _page_role(charts, tables, pictures, texts),
            "text_frames": texts,
            "pictures": pictures,
            "charts": charts,
            "tables": tables,
            "title": title,
        })
    chart_pages = [s["page"] for s in slides if s["charts"] >= 1]
    return {
        "source": pptx_path.name,
        "slide_count": len(slides),
        "slides": slides,
        "rhythm": [s["role"] for s in slides],
        "font_tiers": stats["font_tiers"],
        "font_pts": stats["font_pts"],
        "palette": stats["palette"],
        "scheme_refs": stats["scheme_refs"],
        "chart_frequency": {
            "count": len(chart_pages),
            "pages": chart_pages,
        },
    }


def _sorted(counter: dict) -> list[dict]:
    # (-count, key) ordering keeps output stable across runs/machines.
    items = [{"key": k, "count": v["count"], "pages": v["pages"]}
             for k, v in counter.items()]
    return sorted(items, key=lambda item: (-item["count"], item["key"]))


def _pct(part: int, whole: int) -> str:
    return "0%" if whole == 0 else f"{round(part * 100 / whole)}%"


def _support(pages: list[int]) -> str:
    if len(pages) >= 2:
        return LABEL_MULTI_PAGE
    return LABEL_SINGLE_PAGE


def build_observations(data: dict) -> list[dict]:
    """Derive observations strictly from extracted page evidence.

    Every observation carries its evidence pages; nothing here may summarise
    beyond what the per-page records show (anti-fabrication contract).
    """
    slides = data["slides"]
    total = len(slides)
    if total == 0:
        return []
    observations: list[dict] = []

    def add(text: str, pages: list[int]) -> None:
        if not pages:
            return  # defensive: never emit an observation without provenance
        observations.append({"text": text, "pages": sorted(pages)})

    add(f"开篇页角色为 {slides[0]['role']}", [slides[0]["page"]])
    add(f"收束页角色为 {slides[-1]['role']}", [slides[-1]["page"]])

    role_pages: dict[str, list[int]] = {}
    for slide in slides:
        role_pages.setdefault(slide["role"], []).append(slide["page"])
    for role in sorted(role_pages):
        pages = role_pages[role]
        add(f"角色「{role}」出现 {len(pages)} 页", pages)

    chart_freq = data["chart_frequency"]
    if chart_freq["pages"]:
        add(
            f"图表页 {len(chart_freq['pages'])}/{total}（占比 "
            f"{_pct(len(chart_freq['pages']), total)}），量化论证由原生图表承载",
            chart_freq["pages"],
        )
    chart_free = [s["page"] for s in slides if s["charts"] == 0]
    if chart_free and chart_freq["pages"]:
        add("其余页无原生图表，论证以文本/图像推进", chart_free)

    titled = [s for s in slides if s["title"]]
    add("标题形态：句末带终结标点（。．!？?）的完整句断言标题",
        [s["page"] for s in titled
         if s["title"] and s["title"][-1] in "。．!？?"])
    add("标题形态：无句末标点的短语式标题",
        [s["page"] for s in titled
         if s["title"] and s["title"][-1] not in "。．!？?"])
    add("标题含阿拉伯数字（数字举证式标题）",
        [s["page"] for s in titled if s["title"] and any(c.isdigit() for c in s["title"])])

    for entry in _sorted(data["palette"])[:PALETTE_TOP_N]:
        add(f"显式色 #{entry['key']} 出现 {entry['count']} 次", entry["pages"])
    scheme = data["scheme_refs"]
    if scheme:
        all_scheme_pages = sorted({p for v in scheme.values() for p in v["pages"]})
        add(f"主题 scheme 色引用 {sum(v['count'] for v in scheme.values())} 次"
            f"（具体色值取决于主题，未展开为 HEX）", all_scheme_pages)

    for tier in TIER_ORDER:
        entry = data["font_tiers"].get(tier)
        if entry:
            label = {"h1": "大标层级（≥28pt）", "h2": "中标层级（18–28pt）",
                     "body": "正文层级（14–18pt）", "caption": "注脚层级（<14pt）",
                     FONT_TIER_UNSPECIFIED: "未显式设置字号（继承母版）"}[tier]
            add(f"字号分档「{label}」出现 {entry['count']} 个 run", entry["pages"])
    return observations


def render_observations(observations: list[dict]) -> list[str]:
    """Render observation lines, dropping any entry without provenance.

    This filter is the anti-fabrication hard rule: an observation that lost
    its page evidence must never reach the archive, even if a caller built
    it incorrectly upstream.
    """
    lines = []
    for obs in observations:
        pages = sorted(set(obs.get("pages", [])))
        if not pages:
            continue
        refs = ", ".join(f"p{p}" for p in pages)
        lines.append(f"- {obs['text']}（{refs}）〔{_support(pages)}〕")
    return lines


def _title_sample_line(slide: dict) -> str:
    if not slide["title"]:
        return f"- p{slide['page']}〔{slide['role']}〕（无标题占位符文本）"
    title = slide["title"]
    return f"- p{slide['page']}〔{slide['role']}〕「{title}」（{len(title)} 字）"


def render_markdown(data: dict, observations: list[dict]) -> str:
    slides = data["slides"]
    total = data["slide_count"]
    titled = [s for s in slides if s["title"]]
    chart_freq = data["chart_frequency"]
    lines: list[str] = []
    lines.append(f"# 标杆 deck 蒸馏档案 — {data['source']}")
    lines.append("")
    lines.append(f"> 来源：{data['source']}｜共 {total} 页｜"
                 f"scripts/distill_deck_style.py 确定性提取（无 LLM 概括）")
    lines.append("")
    lines.append("## 一、论证模式观察")
    lines.append("")
    obs_lines = render_observations(observations)
    lines.extend(obs_lines or ["- （无可机读证据支撑的观察）"])
    lines.append("")
    lines.append("## 二、页面节奏（角色序列）")
    lines.append("")
    lines.append("| 页 | 角色 | 文本框 | 图片 | 图表 | 表格 |")
    lines.append("| --- | --- | --- | --- | --- | --- |")
    for s in slides:
        lines.append(f"| p{s['page']} | {s['role']} | {s['text_frames']} | "
                     f"{s['pictures']} | {s['charts']} | {s['tables']} |")
    lines.append("")
    rhythm_text = " → ".join(f"p{s['page']}:{s['role']}" for s in slides)
    lines.append(f"角色序列：{rhythm_text}")
    lines.append("")
    lines.append("## 三、版式偏好")
    lines.append("")
    lines.extend(_layout_preference_lines(slides, chart_freq, total))
    lines.append("")
    lines.append("## 四、标题风格")
    lines.append("")
    lines.append(f"> 标题覆盖率：{len(titled)}/{total} 页（有标题占位符文本）")
    lines.append("")
    lines.extend(_title_sample_line(s) for s in slides)
    lines.append("")
    lines.append("## 五、样本局限声明")
    lines.append("")
    lines.append(f"- 本档案仅基于该 deck {total} 页的机读证据（shape/字号/显式色），"
                 "不含其创作语境、受众与迭代史；反演时属「需确认」的默认组。")
    lines.append(f"- 标题样本覆盖 {len(titled)}/{total} 页；未显式设置字号的 run 计入"
                 "「继承母版」档，不代表实际渲染字号。")
    lines.append("- 以上每条观察均带页码出处；超出采样范围的概括（如动机、团队流程）"
                 "本档案不涉及，引用方不得外推。")
    lines.append("- 反演红线：不得把通用版式（如「一页一要点」）包装成该 deck 独有；"
                 "标注三组——应延续/需确认/偶然成立——见 references/deck-distillation.md。")
    lines.append("")
    return "\n".join(lines)


def _role_counts(slides: list[dict]) -> dict:
    counter: dict[str, dict] = {}
    for slide in slides:
        entry = counter.setdefault(slide["role"], {"count": 0, "pages": []})
        entry["count"] += 1
        entry["pages"].append(slide["page"])
    return counter


def _layout_preference_lines(slides: list[dict], chart_freq: dict, total: int) -> list[str]:
    lines = []
    for role_entry in _sorted(_role_counts(slides)):
        pages = role_entry["pages"]
        members = [s for s in slides if s["role"] == role_entry["key"]]
        tf_range = (min(s["text_frames"] for s in members),
                    max(s["text_frames"] for s in members))
        lines.append(
            f"- 角色「{role_entry['key']}」{len(pages)} 页（{', '.join(f'p{p}' for p in pages)}）："
            f"文本框 {tf_range[0]}–{tf_range[1]} 个"
            + (f"；含原生图表" if role_entry["key"] == "chart" else "")
            + f"〔{_support(pages)}〕")
    if chart_freq["pages"]:
        lines.append(f"- 原生图表频率：{chart_freq['count']}/{total}"
                     f"（{_pct(chart_freq['count'], total)}），页码见第一层观察")
    else:
        lines.append(f"- 全 deck 无原生图表（0/{total}）——量化论证若有，为非图表形态")
    return lines


def render_json(data: dict, observations: list[dict]) -> str:
    payload = {
        "source": data["source"],
        "slide_count": data["slide_count"],
        "slides": data["slides"],
        "rhythm": data["rhythm"],
        "palette": _sorted(data["palette"])[:PALETTE_TOP_N],
        "scheme_color_refs": _sorted(data["scheme_refs"]),
        "font_tiers": [
            {"tier": tier, **data["font_tiers"][tier]}
            for tier in TIER_ORDER if tier in data["font_tiers"]
        ],
        "chart_frequency": data["chart_frequency"],
        "observations": [o for o in observations if o.get("pages")],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Distill a trusted benchmark PPTX into a five-layer "
                    "provenance-cited style archive (caller owns Gate 0 trust).")
    parser.add_argument("pptx", help="Path to a user-vetted trusted PPTX "
                                     "(this script performs no trust decision)")
    parser.add_argument("--json", nargs="?", const="-", default=None, metavar="|-",
                        help="Emit deterministic JSON to stdout instead of markdown")
    parser.add_argument("--out", type=Path, default=None,
                        help="Write the markdown archive to this path "
                             "(parent directory must exist; convention: "
                             "${LEO_PPT_HOME}/decks-styles/<name>.md)")
    args = parser.parse_args(argv)

    pptx_path = Path(args.pptx)
    if not pptx_path.is_file():
        print(f"error: PPTX not found or not a file: {pptx_path}", file=sys.stderr)
        return EXIT_USAGE
    try:
        data = extract(pptx_path)
    except Exception as exc:  # unreadable zip / non-OOXML / corrupt package
        print(f"error: cannot parse PPTX ({type(exc).__name__}: {exc})",
              file=sys.stderr)
        return EXIT_USAGE
    if data["slide_count"] == 0:
        print("error: deck has 0 slides — nothing to distil", file=sys.stderr)
        return EXIT_USAGE

    observations = build_observations(data)
    if args.json is not None:
        print(render_json(data, observations))
    else:
        markdown = render_markdown(data, observations)
        if args.out is not None:
            try:
                args.out.write_text(markdown, encoding="utf-8")
            except OSError as exc:
                print(f"error: cannot write --out ({exc})", file=sys.stderr)
                return EXIT_USAGE
        print(markdown)
    return 0


if __name__ == "__main__":
    sys.exit(main())
