#!/usr/bin/env python3
"""Export a speaker script (讲稿) from a delivered PPTX or a deck master.

Sources (exactly one required):
  --pptx    read per-slide notes via python-pptx (authoritative post-assembly
            source; this is what the audience-facing deck actually carries);
  --master  read ``- speaker_script：`` bullets per ``## P<N>`` page block from
            a deck-master markdown file (pre-assembly source).

Output is deterministic markdown. Pages without notes are listed honestly as
（无备注）and counted in a footer summary — the exporter never invents script
text. Exit codes: 0 = exported; 2 = usage or IO error.

``--prose-check`` runs the speaker-script discipline scan (R-16, reusing
``check_deck_prose.scan_speaker_scripts``: clichés / >40-char sentences /
notice-tone 「大家」) before export and prints a ``PROSE-WARN`` summary to
stderr — advisory only, never blocks the export or changes the exit code.

``--duration-check <page-content-pack.json>`` (R-82) prints a
``DURATION-WARN`` summary to stderr: per-page budgeted seconds (page-level
``budget_seconds`` first, then explicit deck ``duration_seconds`` split
evenly, else unknown) versus the speaking-time estimate of the exported
script text (chars / SPEAKING_CHARS_PER_MINUTE). Advisory only — never
blocks export or changes the exit code.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

SPEAKER_LINE = re.compile(r"^\s*[-*]?\s*speaker_script\s*[:：]\s*(.+)$")
PAGE_HEADING = re.compile(r"^##\s+P(\d+)\s*(.*)$")


def fail(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(2)


def notes_from_pptx(path: Path) -> "list[tuple[str, str]]":
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # pptx raises assorted errors on bad files
        fail(f"cannot open pptx {path}: {exc}")
    pages: list[tuple[str, str]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.text_frame:
            title = slide.shapes.title.text_frame.text.strip().splitlines()[0] if slide.shapes.title.text_frame.text.strip() else ""
        text = ""
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
        pages.append((f"第 {index} 页{('：' + title) if title else ''}", text))
    return pages


def notes_from_master(path: Path) -> "list[tuple[str, str]]":
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as exc:
        fail(f"cannot read master {path}: {exc}")
    pages: list[tuple[str, str]] = []
    current: "tuple[str, list[str]] | None" = None
    for raw in text.splitlines():
        heading = PAGE_HEADING.match(raw)
        if heading:
            if current is not None:
                pages.append((current[0], "\n".join(current[1]).strip()))
            number, title = heading.group(1), heading.group(2).strip()
            current = (f"第 {number} 页{('：' + title) if title else ''}", [])
            continue
        if current is None:
            continue
        line = SPEAKER_LINE.match(raw)
        if line:
            current[1].append(line.group(1).strip())
    if current is not None:
        pages.append((current[0], "\n".join(current[1]).strip()))
    if not pages:
        fail(f"no '## P<N>' page headings found in {path}")
    return pages


def render(source_label: str, pages: "list[tuple[str, str]]") -> str:
    lines = [f"# 讲稿（来源：{source_label}）", ""]
    missing = 0
    for label, text in pages:
        lines.append(f"## {label}")
        lines.append("")
        if text:
            lines.append(text)
        else:
            lines.append("（无备注）")
            missing += 1
        lines.append("")
    lines.append(f"共 {len(pages)} 页；有备注 {len(pages) - missing} 页，缺备注 {missing} 页（缺页请回母版补 speaker_script 后重建 notes）。")
    return "\n".join(lines).rstrip() + "\n"


def _print_prose_warnings(pages: "list[tuple[str, str]]") -> None:
    """Advisory R-16 scan before export; never blocks or changes exit codes."""
    scripts_dir = Path(__file__).resolve().parent
    if str(scripts_dir) not in sys.path:
        sys.path.insert(0, str(scripts_dir))
    try:
        from check_deck_prose import scan_speaker_scripts
    except ImportError as exc:  # keep export working even if sibling moves
        print(f"PROSE-CHECK: 无法加载 check_deck_prose（{exc}），跳过讲稿纪律扫描",
              file=sys.stderr)
        return
    warns = [f for f in scan_speaker_scripts(pages) if f["severity"] == "WARN"]
    if not warns:
        return
    print(f"PROSE-WARN: 讲稿纪律 {len(warns)} 处（不阻断导出；判读见 "
          "deck-master.md 确定性检测小节）", file=sys.stderr)
    for finding in warns:
        print(f"PROSE-WARN: {finding['page']} — {finding['message']}", file=sys.stderr)


def _print_duration_warnings(
    pages: "list[tuple[str, str]]",
    pack_path: Path,
) -> None:
    """Advisory R-82 check; same non-blocking pattern as --prose-check."""

    scripts_dir = Path(__file__).resolve().parent
    runtime_src = scripts_dir.parent / "runtime" / "src"
    for candidate in (runtime_src, scripts_dir):
        if str(candidate) not in sys.path:
            sys.path.insert(0, str(candidate))
    try:
        from leo_ppt_generator.delivery_disclosure import speaker_duration_report
    except ImportError as exc:
        print(f"DURATION-CHECK: 无法加载 delivery_disclosure（{exc}），跳过时长校准",
              file=sys.stderr)
        return
    try:
        pack = json.loads(pack_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        print(f"DURATION-CHECK: 内容包不可读（{exc}），跳过时长校准", file=sys.stderr)
        return
    deck_seconds = (pack.get("deck") or {}).get("duration_seconds")
    budgets = {str(page.get("page_id")): page.get("budget_seconds")
               for page in pack.get("pages", [])}
    # PPTX/master 页序与内容包页序一致（同一母版派生），按序取 page_id。
    page_ids = [str(page.get("page_id")) for page in pack.get("pages", [])]
    rows = []
    for index, (_label, text) in enumerate(pages):
        page_id = page_ids[index] if index < len(page_ids) else f"page_{index + 1:03d}"
        rows.append({"page_id": page_id, "chars": len(text.replace("\n", "").replace(" ", "")),
                     "budget_seconds": budgets.get(page_id)})
    report = speaker_duration_report(pages=rows, deck_duration_seconds=deck_seconds)
    if report["pages_over"]:
        print(f"DURATION-WARN: {report['pages_over']}/{report['pages_total']} 页讲稿估算超预算"
              "（advisory 不阻断；语速模型 "
              f"{report['chars_per_minute']} 字/分）", file=sys.stderr)
        for row in report["pages"]:
            if row["status"] == "over":
                print(f"DURATION-WARN: {row['page_id']} 估算 {row['estimated_seconds']}s "
                      f"> 预算 {row['budget_seconds']}s", file=sys.stderr)
    if report["pages_unknown_budget"]:
        print(f"DURATION-CHECK: {report['pages_unknown_budget']} 页无预算（页级与整册双缺失，"
              "unknown 如实披露）", file=sys.stderr)


def main(argv: "list[str] | None" = None) -> int:
    parser = argparse.ArgumentParser(description="Export a speaker script from PPTX notes or a deck master.")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--pptx", help="delivered PPTX file")
    source.add_argument("--master", help="deck-master markdown file")
    parser.add_argument("--out", help="output markdown path (default stdout)")
    parser.add_argument("--prose-check", action="store_true",
                        help="run the R-16 speaker-script discipline scan before "
                             "export; print a WARN summary to stderr (advisory)")
    parser.add_argument("--duration-check", metavar="PACK",
                        help="page-content-pack.json：R-82 讲稿时长校准（页级预算优先，"
                             "整册均分兜底，双缺失 unknown；advisory）")
    args = parser.parse_args(argv)

    if args.pptx:
        pages = notes_from_pptx(Path(args.pptx))
        label = f"PPTX notes — {Path(args.pptx).name}"
    else:
        pages = notes_from_master(Path(args.master))
        label = f"deck master — {Path(args.master).name}"

    if args.prose_check:
        _print_prose_warnings(pages)

    if args.duration_check:
        _print_duration_warnings(pages, Path(args.duration_check))

    output = render(label, pages)
    if args.out:
        try:
            Path(args.out).write_text(output, encoding="utf-8")
        except OSError as exc:
            fail(f"cannot write {args.out}: {exc}")
    else:
        sys.stdout.write(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
