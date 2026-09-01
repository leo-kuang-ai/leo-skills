#!/usr/bin/env python3
"""distill_deck_style.py 单测：每条观察带页码出处 / 空 deck 与坏输入 exit 2 /
确定性 / JSON 完备 / 局限声明在场 / 无出处观察被拒的构造路径 / 角色推断 /
字号分档 / 色板直方图 / 图表频率 / --out 落盘 / 验证状态标注。"""
import base64
import io
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPT = SKILL_DIR / "scripts" / "distill_deck_style.py"
sys.path.insert(0, str(SKILL_DIR / "scripts"))

import distill_deck_style  # noqa: E402

# Minimal 1x1 transparent PNG for add_picture.
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNg"
    "YGBgAAAABQABh6FO1AAAAABJRU5ErkJggg=="
)


def _fixture_deck(path: Path) -> None:
    """4-page deck: sparse-text / dense-text / chart / text-image."""
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "AI 时代的测试 deck。"
    s1.placeholders[1].text = "副标题：蒸馏演示"
    for run in s1.shapes.title.text_frame.paragraphs[0].runs:
        run.font.size = Pt(40)
        run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    for i, text in enumerate(["要点一", "要点二", "要点三", "要点四"]):
        box = s2.shapes.add_textbox(Inches(0.5 + i), Inches(1), 2, 1)
        box.text_frame.text = text
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(20 if i < 2 else 16)
        run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("s", (1, 2))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                        Inches(1), Inches(1), 4, 3, chart_data)
    box = s3.shapes.add_textbox(Inches(0.3), Inches(0.3), 3, 0.5)
    box.text_frame.text = "增长 3 倍的论证页"
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4.shapes.add_picture(io.BytesIO(_PNG), Inches(1), Inches(1), 2, 2)
    box = s4.shapes.add_textbox(Inches(0.3), Inches(0.3), 3, 0.5)
    box.text_frame.text = "案例图解"
    prs.save(str(path))


def _run(args, expect_ok=True):
    result = subprocess.run(
        [sys.executable, str(SCRIPT)] + [str(a) for a in args],
        capture_output=True, text=True,
    )
    if expect_ok and result.returncode != 0:
        raise AssertionError(
            f"expected exit 0, got {result.returncode}: {result.stderr}")
    return result


class FixtureDeckTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pptx = Path(self._tmp.name) / "deck.pptx"
        _fixture_deck(self.pptx)
        self.addCleanup(self._tmp.cleanup)
        self.md = _run([self.pptx]).stdout

    def _layer_one_lines(self):
        body = self.md.split("## 一、论证模式观察", 1)[1]
        body = body.split("## 二、", 1)[0]
        return [line for line in body.splitlines() if line.startswith("- ")]

    def test_every_observation_carries_page_provenance(self):
        lines = self._layer_one_lines()
        self.assertGreaterEqual(len(lines), 8)
        for line in lines:
            self.assertRegex(line, r"（p\d+(, p\d+)*）"), line

    def test_deterministic_output_across_runs(self):
        second = _run([self.pptx]).stdout
        self.assertEqual(self.md, second)

    def test_rhythm_sequence_matches_fixture(self):
        self.assertIn(
            "角色序列：p1:sparse-text → p2:dense-text → p3:chart → p4:text-image",
            self.md)

    def test_layout_role_inference_chart_page(self):
        self.assertRegex(self.md, r"\| p3 \| chart \| 1 \| 0 \| 1 \| 0 \|")

    def test_font_tier_buckets_present(self):
        self.assertIn("字号分档「大标层级（≥28pt）」出现 1 个 run", self.md)
        self.assertIn("字号分档「中标层级（18–28pt）」出现 2 个 run", self.md)
        self.assertIn("字号分档「正文层级（14–18pt）」出现 2 个 run", self.md)

    def test_palette_histogram_with_pages(self):
        self.assertIn("显式色 #C0392B 出现 5 次（p1, p2）", self.md)

    def test_chart_frequency_observation(self):
        self.assertIn("图表页 1/4（占比 25%）", self.md)

    def test_limitation_statement_present(self):
        body = self.md.split("## 五、样本局限声明", 1)[1]
        self.assertIn("仅基于该 deck 4 页的机读证据", body)
        self.assertIn("不得外推", body)
        self.assertIn("不得把通用版式", body)

    def test_support_labels_single_vs_multi_page(self):
        lines = self._layer_one_lines()
        self.assertTrue(any("（p1, p2, p4）〔多页复现〕" in line for line in lines))
        self.assertTrue(
            any("〔单页观察，反演时归「需确认」组〕" in line for line in lines))

    def test_title_style_layer_lists_samples_with_pages(self):
        self.assertIn("标题覆盖率：1/4 页", self.md)
        self.assertIn("p1〔sparse-text〕「AI 时代的测试 deck。」（14 字）", self.md)
        self.assertIn("p2〔dense-text〕（无标题占位符文本）", self.md)


class JsonOutputTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pptx = Path(self._tmp.name) / "deck.pptx"
        _fixture_deck(self.pptx)
        self.addCleanup(self._tmp.cleanup)

    def test_json_payload_complete(self):
        payload = json.loads(_run([self.pptx, "--json"]).stdout)
        for key in ("source", "slide_count", "slides", "rhythm", "palette",
                    "scheme_color_refs", "font_tiers", "chart_frequency",
                    "observations"):
            self.assertIn(key, payload)
        self.assertEqual(payload["slide_count"], 4)
        self.assertEqual(payload["rhythm"],
                         ["sparse-text", "dense-text", "chart", "text-image"])
        for slide in payload["slides"]:
            for key in ("page", "role", "text_frames", "pictures",
                        "charts", "tables"):
                self.assertIn(key, slide)
        self.assertEqual(payload["chart_frequency"], {"count": 1, "pages": [3]})
        top = payload["palette"][0]
        self.assertEqual(top["key"], "C0392B")
        self.assertEqual(top["pages"], [1, 2])
        for obs in payload["observations"]:
            self.assertTrue(obs["pages"])  # JSON side keeps provenance too

    def test_json_deterministic(self):
        first = _run([self.pptx, "--json"]).stdout
        second = _run([self.pptx, "--json"]).stdout
        self.assertEqual(first, second)

    def test_out_writes_markdown_file(self):
        out = Path(self._tmp.name) / "archive.md"
        stdout_md = _run([self.pptx]).stdout
        _run([self.pptx, "--out", out])
        # stdout = raw markdown + one trailing newline from print().
        self.assertEqual(out.read_text(encoding="utf-8") + "\n", stdout_md)

    def test_out_missing_parent_dir_exits_2(self):
        out = Path(self._tmp.name) / "no-such-dir" / "archive.md"
        result = _run([self.pptx, "--out", out], expect_ok=False)
        self.assertEqual(result.returncode, 2)


class InvalidInputTest(unittest.TestCase):
    def test_missing_file_exits_2(self):
        result = _run([Path("/nonexistent/deck.pptx")], expect_ok=False)
        self.assertEqual(result.returncode, 2)

    def test_non_pptx_file_exits_2(self):
        with tempfile.TemporaryDirectory() as tmp:
            bogus = Path(tmp) / "deck.pptx"
            bogus.write_text("not a zip", encoding="utf-8")
            result = _run([bogus], expect_ok=False)
        self.assertEqual(result.returncode, 2)

    def test_empty_deck_exits_2(self):
        with tempfile.TemporaryDirectory() as tmp:
            empty = Path(tmp) / "empty.pptx"
            Presentation().save(str(empty))
            result = _run([empty], expect_ok=False)
        self.assertEqual(result.returncode, 2)
        self.assertIn("0 slides", result.stderr)


class AntiFabricationTest(unittest.TestCase):
    def test_unprovenanced_observation_dropped_at_render(self):
        lines = distill_deck_style.render_observations([
            {"text": "有出处的观察", "pages": [2]},
            {"text": "无出处的概括（必须被丢弃）", "pages": []},
            {"text": "缺 pages 键的观察（必须被丢弃）"},
        ])
        self.assertEqual(len(lines), 1)
        self.assertIn("（p2）", lines[0])
        self.assertNotIn("无出处", "\n".join(lines))
        self.assertNotIn("缺 pages", "\n".join(lines))

    def test_built_observations_never_empty_pages(self):
        with tempfile.TemporaryDirectory() as tmp:
            pptx = Path(tmp) / "deck.pptx"
            _fixture_deck(pptx)
            data = distill_deck_style.extract(pptx)
        for obs in distill_deck_style.build_observations(data):
            self.assertTrue(obs["pages"])


if __name__ == "__main__":
    unittest.main()
