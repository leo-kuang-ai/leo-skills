"""Object-level PPTX builder: manifest IR → python-pptx object API compiler.

This is the leo-owned replacement for the vendored ``write_pptx``/``write_deck``
zip writer (kept as the ``legacy`` builder behind ``LEO_EDITABLE_BUILDER``).
Interface shape is identical to legacy so the adapter can dispatch with zero
caller changes:

- ``write_pptx(manifest, out_path, manifest_path)``
- ``write_deck(deck, page_entries, out_path, notes_entries)``
- ``normalize_manifest(manifest)`` (delegates to the shared geometry layer)
- ``TEXT_ALIGNMENTS`` / ``TEXT_VERTICAL_ALIGNMENTS`` (re-exported constants)

Differences from legacy are intentional and listed in the team design
(§3.1.7): default python-pptx template parts (11 layouts, full master/theme),
endParaRPr, and ``docProps/app.xml`` identity ``leo-ppt-generator/<id>``.
Equivalence is defined on the structural projection
(``object_projection.py``), never on bytes.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.parts.image import ImagePart
from pptx.util import Emu

from . import _oxml
from .deterministic_zip import save_canonical
from . import geometry

BUILDER_ID = "leo-ppt-generator/object-builder-1"
DEFAULT_FONT = "PingFang SC"
BLANK_LAYOUT_INDEX = 6

TEXT_ALIGNMENTS = geometry.TEXT_ALIGNMENTS
TEXT_VERTICAL_ALIGNMENTS = geometry.TEXT_VERTICAL_ALIGNMENTS

ALIGN_TO_PARA = {
    "l": PP_ALIGN.LEFT,
    "ctr": PP_ALIGN.CENTER,
    "r": PP_ALIGN.RIGHT,
}


def normalize_manifest(manifest: dict) -> dict:
    """Shared-geometry normalization (semantics identical to legacy)."""
    return geometry.normalize_manifest(manifest)


class _CompileContext:
    """Per-deck mutable state: media numbering, ids, notes freeze bookkeeping."""

    def __init__(self) -> None:
        self.media_index = 1
        self.notes_parts: dict[int, object] = {}
        self.notes_replacements: dict[str, bytes] = {}

    def next_media(self) -> int:
        index = self.media_index
        self.media_index += 1
        return index


def _resolve_image_path(item: dict, manifest_path) -> Path:
    src = Path(item["path"])
    if not src.is_absolute():
        src = Path(manifest_path).resolve().parent / src
    return src


def _image_ext(path) -> str:
    suffix = Path(path).suffix.lower()
    return ".jpg" if suffix == ".jpeg" else suffix


def _content_type_for(path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix == ".png":
        return "image/png"
    if suffix in (".jpg", ".jpeg"):
        return "image/jpeg"
    if suffix == ".gif":
        return "image/gif"
    if suffix == ".svg":
        return "image/svg+xml"
    raise ValueError(f"Unsupported image type: {path}")


def _compile_image(slide, item: dict, manifest_path, ctx: _CompileContext, shape_id: int) -> None:
    """One media part per manifest reference, named ``image{N}{ext}``.

    python-pptx's package-level sha1 dedup would collapse repeated references
    into a single part and renumber media, which breaks the validator's
    ``image{index}`` reverse lookup and its media-count contract; SVG is not
    supported by ``pptx.image.Image`` at all. Parts are therefore constructed
    directly and related to the slide part here.
    """
    src = _resolve_image_path(item, manifest_path)
    blob = src.read_bytes()
    partname = PackURI(f"/ppt/media/image{ctx.next_media()}{_image_ext(src)}")
    part = ImagePart(partname, _content_type_for(src), slide.part.package, blob, src.name)
    rel_id = slide.part.relate_to(part, RT.IMAGE)
    name = item.get("alt") or Path(item.get("path", "")).stem or f"Image {shape_id}"
    pic = _oxml.build_pic_element(
        shape_id,
        name,
        rel_id,
        geometry.emu(item.get("left", 0)),
        geometry.emu(item.get("top", 0)),
        geometry.emu(item.get("width", 1)),
        geometry.emu(item.get("height", 1)),
    )
    slide.shapes._spTree.append(pic)


def _compile_textbox(slide, item: dict, shape_id: int, theme_fonts: dict | None) -> None:
    textbox = slide.shapes.add_textbox(
        Emu(geometry.emu(item.get("left", 0))),
        Emu(geometry.emu(item.get("top", 0))),
        Emu(geometry.emu(item.get("width", 1))),
        Emu(geometry.emu(item.get("height", 0.4))),
    )
    _oxml.set_cnvpr_id_name(textbox._element, shape_id, f"TextBox {shape_id}")
    rotation = item.get("rotation")
    if rotation not in (None, ""):
        textbox.rotation = float(rotation)

    wrap = item.get("wrap", "none")
    anchor = geometry.text_vertical_alignment(item.get("valign", "top"))[0]
    autofit_element = "spAutoFit" if item.get("autofit") == "shape" else "noAutofit"
    _oxml.configure_textbox_bodyPr(textbox.text_frame, wrap, anchor, autofit_element)
    # legacy textboxes carry an explicit noFill outline
    textbox.line.fill.background()

    align_token = geometry.text_alignment(item.get("align", "left"))[0]
    paragraphs = item.get("paragraphs")
    runs = item.get("runs")
    if paragraphs:
        paragraph_payloads = [
            [{"text": paragraph}] if isinstance(paragraph, str) else paragraph.get("runs", [{"text": paragraph.get("text", "")}])
            for paragraph in paragraphs
        ]
    elif runs:
        paragraph_payloads = [runs]
    else:
        paragraph_payloads = [[{"text": part}] for part in str(item.get("text", "")).split("\n")]

    text_frame = textbox.text_frame
    for position, payload in enumerate(paragraph_payloads):
        paragraph = text_frame.paragraphs[0] if position == 0 else text_frame.add_paragraph()
        paragraph.alignment = ALIGN_TO_PARA[align_token]
        for run_spec in payload:
            run = paragraph.add_run()
            _apply_run(run, run_spec, item, theme_fonts)


def _resolve_font(run_spec: dict, item: dict, theme_fonts: dict | None) -> str:
    """Run → box → theme slot (by role) → legacy default, in that priority."""
    if run_spec.get("font"):
        return str(run_spec["font"])
    if item.get("font"):
        return str(item["font"])
    if theme_fonts:
        role = str(item.get("role", "body")).strip().lower()
        slot = theme_fonts.get("head_font_face") if role == "title" else theme_fonts.get("body_font_face")
        if slot:
            return str(slot)
    return DEFAULT_FONT


def _apply_run(run, run_spec: dict, item: dict, theme_fonts: dict | None) -> None:
    run.text = str(run_spec.get("text", ""))
    font_size = run_spec.get("font_size", item.get("font_size", 18))
    # legacy truncating centipoint math (Pt() rounds and would differ)
    _oxml.set_run_sz(run, font_size)
    _oxml.set_run_font_slots(run, _resolve_font(run_spec, item, theme_fonts))
    color = geometry.hex_color(run_spec.get("color", item.get("color", "#111111")))
    run.font.color.rgb = RGBColor.from_string(color)
    if run_spec.get("bold", item.get("bold")):
        run.font.bold = True
    if run_spec.get("italic", item.get("italic")):
        run.font.italic = True
    baseline = run_spec.get("baseline")
    if baseline not in (None, ""):
        _oxml.set_run_baseline(run, baseline)


def _apply_shape_fill_and_line(shape, item: dict) -> None:
    fill = item.get("fill")
    if hasattr(shape, "fill"):
        if not fill or fill == "none":
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(geometry.hex_color(fill))
    else:
        # connectors have no FillFormat surface
        _oxml.set_connector_fill(shape._element, geometry.hex_color(fill) if fill and fill != "none" else None)
    stroke = item.get("stroke", "#000000")
    if not stroke or stroke == "none":
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(geometry.hex_color(stroke))
        # legacy: w = int(float(stroke_width or 1) * 12700) — 0 falls back to 1
        shape.line.width = Emu(int(float(item.get("stroke_width", 1) or 1) * 12700))
        dash = item.get("dash")
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE

            try:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.from_xml(str(dash))
            except ValueError as exc:
                raise ValueError(f"Unsupported line dash value: {dash}") from exc


def _compile_shape(slide, item: dict, shape_id: int) -> None:
    kind = item.get("type", "rect")
    polygon_px = item.get("polygon_px")
    preset = item.get("preset")
    if polygon_px:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(geometry.emu(item.get("left", 0))),
            Emu(geometry.emu(item.get("top", 0))),
            Emu(geometry.emu(item.get("width", 1))),
            Emu(geometry.emu(item.get("height", 1))),
        )
        _oxml.replace_geometry_with_custgeom(shape, polygon_px, item.get("box_px"))
    elif kind == "line" and not preset:
        points = item.get("points")
        if points and len(points) == 4:
            begin = (Emu(geometry.emu(points[0])), Emu(geometry.emu(points[1])))
            end = (Emu(geometry.emu(points[2])), Emu(geometry.emu(points[3])))
        else:
            begin = (Emu(geometry.emu(item.get("left", 0))), Emu(geometry.emu(item.get("top", 0))))
            end = (
                Emu(geometry.emu(item.get("left", 0)) + geometry.emu(item.get("width", 0))),
                Emu(geometry.emu(item.get("top", 0)) + geometry.emu(item.get("height", 0))),
            )
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *begin, *end)
    else:
        if not preset:
            preset = (
                "line" if kind == "line" else "ellipse" if kind == "ellipse" else "roundRect" if kind == "roundRect" else "rect"
            )
        try:
            mso_shape = MSO_SHAPE.from_xml(str(preset))
        except ValueError as exc:
            # legacy silently wrote illegal prstGeom tokens; fail at build time instead
            raise ValueError(f"Unsupported shape preset: {preset}") from exc
        shape = slide.shapes.add_shape(
            mso_shape,
            Emu(geometry.emu(item.get("left", 0))),
            Emu(geometry.emu(item.get("top", 0))),
            Emu(geometry.emu(item.get("width", 1))),
            Emu(geometry.emu(item.get("height", 1))),
        )
        if preset == "roundRect":
            adjustment = geometry.round_rect_adjustment(item)
            if adjustment is not None:
                shape.adjustments[0] = adjustment / 100000.0
    _oxml.set_cnvpr_id_name(shape._element, shape_id, f"{str(kind).title()} {shape_id}")
    rotation = item.get("rotation")
    if rotation not in (None, ""):
        shape.rotation = float(rotation)
    if item.get("flip_h") or item.get("flip_v"):
        _oxml.set_xfrm_flip(shape._element, bool(item.get("flip_h")), bool(item.get("flip_v")))
    _apply_shape_fill_and_line(shape, item)


def _compile_table(slide, item: dict, shape_id: int) -> None:
    """Native ``tables[]`` object face (F1-T5, optional manifest section)."""
    rows = item.get("rows") or []
    row_count = len(rows)
    column_count = max((len(row) for row in rows), default=0)
    if row_count == 0 or column_count == 0:
        raise ValueError("tables[] requires a non-empty rows grid")
    frame = slide.shapes.add_table(
        row_count,
        column_count,
        Emu(geometry.emu(item.get("left", 0))),
        Emu(geometry.emu(item.get("top", 0))),
        Emu(geometry.emu(item.get("width", 1))),
        Emu(geometry.emu(item.get("height", 0.5 * row_count))),
    )
    _oxml.set_cnvpr_id_name(frame._element, shape_id, f"Table {shape_id}")
    table = frame.table
    total_width = geometry.emu(item.get("width", 1))
    column_fractions = item.get("col_widths") or [1.0 / column_count] * column_count
    for column, fraction in zip(table.columns, column_fractions):
        column.width = Emu(int(total_width * float(fraction)))
    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            table.cell(row_index, column_index).text = "" if value is None else str(value)
    for merge in item.get("merges") or []:
        table.cell(int(merge[0]), int(merge[1])).merge(table.cell(int(merge[2]), int(merge[3])))


def _compile_background(slide, manifest: dict) -> None:
    background = manifest.get("slide", {}).get("background")
    if not background:
        return
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(geometry.hex_color(background, "FFFFFF"))


def _apply_notes(slide, page_index: int, note: dict | None, ctx: _CompileContext) -> None:
    if not note:
        return
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = str(note.get("text", ""))
    ctx.notes_parts[page_index] = notes_slide.part
    notes_xml = note.get("notes_xml")
    if notes_xml and Path(notes_xml).exists():
        # frozen bytes are re-injected during the canonical repack pass
        ctx.notes_replacements[f"ppt/notesSlides/notesSlide{page_index}.xml"] = Path(notes_xml).read_bytes()


def compile_page(
    slide,
    manifest: dict,
    manifest_path,
    *,
    page_index: int = 1,
    notes: dict | None = None,
    context: _CompileContext | None = None,
    theme_fonts: dict | None = None,
    shape_id_start: int = 2,
) -> None:
    """Compile one normalized manifest onto ``slide`` (page.pptx and deck share this)."""
    ctx = context if context is not None else _CompileContext()
    _compile_background(slide, manifest)

    layered: list[tuple[float, int, str, dict]] = []
    for index, item in enumerate(manifest.get("shapes", [])):
        layered.append((float(item.get("z_index", 100)), index, "shape", item))
    for index, item in enumerate(manifest.get("tables", [])):
        layered.append((float(item.get("z_index", 150)), index, "table", item))
    for rel_index, item in enumerate(manifest.get("images", []), start=1):
        layered.append((float(item.get("z_index", 200)), rel_index, "image", item))
    for index, item in enumerate(manifest.get("text_boxes", [])):
        layered.append((float(item.get("z_index", 300)), index, "text", item))

    shape_id = shape_id_start
    for _z_index, _order, kind, item in sorted(layered, key=lambda entry: (entry[0], entry[1])):
        if kind == "shape":
            _compile_shape(slide, item, shape_id)
        elif kind == "table":
            _compile_table(slide, item, shape_id)
        elif kind == "image":
            _compile_image(slide, item, manifest_path, ctx, shape_id)
        else:
            _compile_textbox(slide, item, shape_id, theme_fonts)
        shape_id += 1
    _apply_notes(slide, page_index, notes, ctx)


def _deck_slide_size(deck: dict, page_entries: list) -> tuple[float, float]:
    slide = deck.get("slide") or {}
    if not slide and page_entries:
        slide = page_entries[0]["manifest"].get("slide", {})
    return float(slide.get("width", 13.333)), float(slide.get("height", 7.5))


def _start_presentation(width_in: float, height_in: float) -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(geometry.emu(width_in))
    prs.slide_height = Emu(geometry.emu(height_in))
    _oxml.set_sldsz_type(prs, width_in, height_in, geometry.is_wide_slide(width_in, height_in))
    return prs


def _finish_package(prs, ctx: _CompileContext, theme_fonts: dict | None, out_path) -> Path:
    # renumber notes slides to their page positions (creation order numbering
    # would detach notesSlide{N} from the validator's page mapping)
    for page_index, part in ctx.notes_parts.items():
        part.partname = PackURI(f"/ppt/notesSlides/notesSlide{page_index}.xml")
    if theme_fonts:
        theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        _oxml.set_theme_font_slots(theme_part, theme_fonts.get("head_font_face"), theme_fonts.get("body_font_face"))
    _oxml.set_app_xml_application(prs.part.package, BUILDER_ID)
    return save_canonical(prs, out_path, replacements=ctx.notes_replacements)


def _manifest_theme(manifest: dict) -> dict | None:
    theme = manifest.get("theme")
    return theme if isinstance(theme, dict) else None


def write_pptx(manifest: dict, out_path, manifest_path) -> Path:
    normalized = normalize_manifest(manifest)
    width_in, height_in = geometry.slide_size(normalized)
    theme_fonts = _manifest_theme(normalized)
    prs = _start_presentation(width_in, height_in)
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT_INDEX])
    ctx = _CompileContext()
    compile_page(slide, normalized, manifest_path, page_index=1, context=ctx, theme_fonts=theme_fonts)
    return _finish_package(prs, ctx, theme_fonts, out_path)


def write_deck(deck: dict, page_entries: list, out_path, notes_entries: list) -> Path:
    if not page_entries:
        raise ValueError("Deck has no pages")
    width_in, height_in = _deck_slide_size(deck, page_entries)
    notes_by_page = {int(entry.get("page_index", 0)): entry for entry in notes_entries if entry.get("text")}
    normalized_entries = [{**entry, "manifest": normalize_manifest(entry["manifest"])} for entry in page_entries]
    theme_fonts = _manifest_theme(deck) or _manifest_theme(normalized_entries[0]["manifest"])
    prs = _start_presentation(width_in, height_in)
    ctx = _CompileContext()
    for slide_index, entry in enumerate(normalized_entries, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT_INDEX])
        compile_page(
            slide,
            entry["manifest"],
            entry["manifest_path"],
            page_index=slide_index,
            notes=notes_by_page.get(slide_index),
            context=ctx,
            theme_fonts=theme_fonts,
        )
    return _finish_package(prs, ctx, theme_fonts, out_path)
