#!/usr/bin/env python3
"""U7/§3：image 完整 deck 端到端（冻结设计 → 真实付费生成 → QA → PPTX → 收据）。

与已完成的 HTML deck（verify_html_deck_e2e.py）消费同一冻结设计：
u19 铜贸复盘 6 页，design_digest 34022cfc…（finance-navy / light）。
生成走生产命令面：``leo-ppt upstream --backend-contract <contract>
codex-ppt -- image generate --size 2560x1440 --prompt-file <p> --out <o>``。

预算纪律（用户 2026-09-08 批准"最低 6 张"）：恰好 6 次 generate 调用、
无样张、无重试加跑（上游 transport 重试不计新张）；脚本对调用次数硬断言。

Usage: runtime/.venv/bin/python scripts/verify_image_deck_e2e.py
"""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "runtime" / "src"))

from leo_ppt_generator.templates import compose_design, project_design_to_prompt  # noqa: E402

LEO = str(SKILL / "runtime" / ".venv" / "bin" / "leo-ppt")
OUT = SKILL / "evals" / "fixtures" / "template-quality" / "image-deck-e2e"

# 与 HTML deck 相同的 6 页冻结输入（verify_html_deck_e2e.PAGES/SLIDE_DATA 摘要）
DESIGN_DIGEST_EXPECTED = ("34022cfc36e038dae2debf121589ea6d45c94123375e689505ebc863e636c3b3")

PAGES = [
    {"page_no": 1, "page_role": "cover", "layout": "builtin:layout:cover-basic",
     "content_ref": "u19-cover"},
    {"page_no": 2, "page_role": "content", "layout": "builtin:layout:body-basic",
     "content_ref": "u19-content-1"},
    {"page_no": 3, "page_role": "data", "layout": "builtin:layout:p25-spec-table",
     "content_ref": "u19-data-1",
     "slots": {"rows": [["敞口对照", "行情 vs 敞口", "结构/执行归因"],
                        ["偏差还原", "事实链", "损失量化口径"],
                        ["制度修订", "权限与阈值", "委员会批准"]],
               "columns": ["议题", "口径", "产出"]}},
    {"page_no": 4, "page_role": "content", "layout": "builtin:layout:body-basic",
     "content_ref": "u19-content-2"},
    {"page_no": 5, "page_role": "data", "layout": "builtin:layout:p25-spec-table",
     "content_ref": "u19-data-2",
     "slots": {"rows": [["触发", "边界条件", "事实链起点"],
                        ["发现", "监控口径", "当日盘后"],
                        ["处置", "权限与阈值", "委员会批准"]],
               "columns": ["环节", "口径", "结论"]}},
    {"page_no": 6, "page_role": "closing", "layout": "builtin:layout:pull-quote",
     "content_ref": "u19-closing"},
]

REQUIRED_TEXT = {
    1: ["铜贸业务风险敞口与对冲策略复盘", "大宗商品贸易 · 风险复盘",
        "结构钱与执行钱 · 偏差事实链 · 制度修订", "风险管理部", "经营班子"],
    2: ["复盘要完成三件事", "季度行情与公司敞口的对照还原（结构钱 vs 执行钱）",
        "两处套保偏差的事实链（触发/发现/处置/影响量化口径）",
        "套保制度修订建议与剩余风险"],
    3: ["复盘结论与产出", "议题", "口径", "产出", "敞口对照", "行情 vs 敞口",
        "结构/执行归因", "偏差还原", "事实链", "损失量化口径", "制度修订",
        "权限与阈值", "委员会批准"],
    4: ["套保偏差事实链", "触发：边界条件触碰的事实链起点",
        "发现：监控口径下的当日盘后识别", "处置：按权限与阈值上报委员会批准"],
    5: ["制度修订要点", "环节", "口径", "结论", "触发", "边界条件", "事实链起点",
        "发现", "监控口径", "当日盘后", "处置", "权限与阈值", "委员会批准"],
    6: ["结构钱要赚得明白，执行钱要亏得清楚。", "复盘结论", "委员会审批口径"],
}

ROLE_LOCK = {
    1: "cover（封面）：主标题可放大；页面留白大气。",
    2: "body（正文要点页）：正文页标题光学尺寸与其他正文页一致。",
    3: "body（数据表格页）：表格表头深藏青底白字，数字右对齐，斑马纹行。",
    4: "body（正文要点页）：正文页标题光学尺寸与其他正文页一致。",
    5: "body（数据表格页）：表格表头深藏青底白字，数字右对齐，斑马纹行。",
    6: "closing（收尾金句页）：金句居中大字，出处小字。",
}

AVOID = [
    "禁止无来源数字进入表格", "涨跌不得只用红绿区分（需符号+单位）", "禁止霓虹强调色",
    "禁止水印、伪造 logo、多余页码、页脚日期戳", "禁止英文填充文本、编造的微小标注、占位符文字",
]


def _style_lock(projection: dict) -> str:
    colors = dict(item.split("=", 1) for item in projection["colors"])
    return (
        f"外层壳锁（逐页一致，不得漂移）：纸底 {colors['background']}，"
        f"主色藏青 {colors['primary']}（标题/表头底），强调 {colors['accent']}，"
        f"正文深灰 {colors['text']}；中文字体 Noto Sans SC（标题 700/正文 400）；"
        "页码固定右下角小字；版面 16:9 全幅；克制、可核验的金融审美，无装饰渐变。"
    )


def _page_prompt(spec: dict, projection: dict) -> str:
    n = spec["page_no"]
    lines = [
        f"为一份中文金融风控复盘演示稿生成第 {n} 页（共 6 页）的整页幻灯片图。",
        f"画幅硬合同：2560x1440（16:9）全幅页面图，不得出现非 16:9 画幅。",
        _style_lock(projection),
        f"页面角色锁：{ROLE_LOCK[n]}",
        "视觉翻译：页面必须完整成稿——以下文字是内容本身，必须逐字原样出现在图上"
        "（含全角标点），不是留白等待配字的装饰：",
    ]
    for text in REQUIRED_TEXT[n]:
        lines.append(f"- 「{text}」")
    lines.append("要求：文字清晰可读、无错字、无重叠；中文排版；层级分明；")
    lines.append("表格页的数字与文字必须严格按上表，不得增删改任何单元格。")
    lines.append("Avoid 清单（生成前逐条自查）：" + "；".join(AVOID) + "。")
    return "\n".join(lines)


def main() -> int:
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "input").mkdir(parents=True)
    (OUT / "prompts").mkdir()
    (OUT / "origin_image").mkdir()

    # 1) 冻结设计（与 HTML deck 同一组合 → 同一 digest）
    from leo_ppt_generator.asset_resolver import AssetResolver
    resolver = AssetResolver(library=SKILL / "template-library")
    design = compose_design("finance-navy", pages=PAGES, resolver=resolver)
    if design["design_digest"] != DESIGN_DIGEST_EXPECTED:
        print(json.dumps({"error": "design_digest_drift",
                          "got": design["design_digest"]}, ensure_ascii=False))
        return 1
    projections = [project_design_to_prompt(design, spec) for spec in PAGES]
    (OUT / "resolved-design.json").write_text(
        json.dumps(design, ensure_ascii=False, indent=2), encoding="utf-8")

    # 2) 冻结后端合同（首选渠道 ark / doubao-seedream-4）
    contract = OUT / "input" / "backend-contract.json"
    proc = subprocess.run([LEO, "backend", "create", "--provider", "ark",
                           "--mode", "generate", "--output", str(contract)],
                          capture_output=True, text=True, timeout=120)
    if proc.returncode != 0:
        print(json.dumps({"error": "backend_create_failed",
                          "stderr": proc.stderr[:200]}, ensure_ascii=False))
        return 1

    # 3) 逐页真实生成（恰好 6 次 generate 调用）
    pages_out, calls = [], 0
    for spec, projection in zip(PAGES, projections):
        n = spec["page_no"]
        prompt_path = OUT / "prompts" / f"slide_{n:02d}.txt"
        prompt_path.write_text(_page_prompt(spec, projection), encoding="utf-8")
        candidate = OUT / "origin_image" / f"slide_{n:02d}.png"
        cmd = [LEO, "upstream", "--backend-contract", str(contract),
               "codex-ppt", "--", "image", "generate",
               "--size", "2560x1440", "--prompt-file", str(prompt_path),
               "--out", str(candidate)]
        gen = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        calls += 1
        entry = {"page_no": n, "role": spec["page_role"],
                 "returncode": gen.returncode,
                 "stderr_tail": gen.stderr.strip()[-200:] if gen.returncode else ""}
        if gen.returncode == 0 and candidate.is_file():
            entry["png"] = str(candidate)
            entry["sha256"] = hashlib.sha256(candidate.read_bytes()).hexdigest()
            entry["bytes"] = candidate.stat().st_size
        pages_out.append(entry)

    # 4) 确定性像素质检（E2 闸门，目录批量）
    qa_report = OUT / "reports" / "visual-qa.json"
    qa_report.parent.mkdir(exist_ok=True)
    qa = subprocess.run([sys.executable, str(SKILL / "scripts" / "visual_qa.py"),
                         str(OUT / "origin_image"), "--report", str(qa_report)],
                        capture_output=True, text=True, timeout=300)

    # 5) PPTX 组装 + 回读（仅全部 6 页生成成功时）
    pptx_ok, pptx_path, slides = False, OUT / "u19-copper-risk-review-image.pptx", 0
    if all(p.get("png") for p in pages_out):
        from pptx import Presentation
        from pptx.util import Emu

        presentation = Presentation()
        presentation.slide_width = Emu(12192000)
        presentation.slide_height = Emu(6858000)
        blank = presentation.slide_layouts[6]
        for page in pages_out:
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(page["png"], 0, 0,
                                     width=presentation.slide_width,
                                     height=presentation.slide_height)
        presentation.save(pptx_path)
        readback = Presentation(pptx_path)
        slides = len(readback.slides)
        pptx_ok = slides == 6 and all(
            len([s for s in slide.shapes if s.shape_type == 13]) == 1
            for slide in readback.slides)

    receipt = {
        "kind": "image-deck-e2e-receipt", "schema_version": 1,
        "authorization": "用户 2026-09-08 批准最低 6 张付费生成",
        "input": "u19 铜贸复盘 6 页（与 HTML deck 同冻结设计）",
        "design_digest": design["design_digest"],
        "design_digest_matches_html_deck": design["design_digest"] == DESIGN_DIGEST_EXPECTED,
        "backend_contract": json.loads(contract.read_text(encoding="utf-8")),
        "generate_calls": calls,
        "budget_discipline": {"authorized": 6, "actual": calls, "ok": calls == 6},
        "pages": pages_out,
        "visual_qa": {"exit": qa.returncode,
                      "report": str(qa_report) if qa_report.exists() else "",
                      "stderr_tail": qa.stderr.strip()[-200:]},
        "pptx": {"path": str(pptx_path), "slides": slides, "readback_ok": pptx_ok,
                 "sha256": hashlib.sha256(pptx_path.read_bytes()).hexdigest()
                 if pptx_path.exists() else ""},
    }
    (OUT / "receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    generated = sum(1 for p in pages_out if p.get("png"))
    ok = (calls == 6 and generated == 6 and qa.returncode == 0 and pptx_ok)
    print(json.dumps({"design_digest": design["design_digest"],
                      "generate_calls": calls, "generated_ok": generated,
                      "visual_qa_exit": qa.returncode,
                      "pptx_readback_ok": pptx_ok, "overall_pass": ok},
                     ensure_ascii=False))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
